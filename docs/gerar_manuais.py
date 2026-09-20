"""Build the printable manuals from the checked-in Markdown sources."""
from pathlib import Path
import re

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def build(source):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.04
    for style, size in (("Title", 23), ("Heading 1", 19), ("Heading 2", 13)):
        document.styles[style].font.name = "Calibri"
        document.styles[style].font.size = Pt(size)
        document.styles[style].font.color.rgb = RGBColor(0, 0, 0)
        document.styles[style].paragraph_format.space_before = Pt(10)
        document.styles[style].paragraph_format.space_after = Pt(7)
    header = section.header.paragraphs[0]
    header.text = "CONFLITOS DE SANGUE"
    header.style = "Caption"
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run("Conflitos de Sangue | ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    first = True
    for line in source.read_text(encoding="utf-8").splitlines():
        if not line:
            continue
        if line == "---":
            document.add_page_break()
        elif line.startswith("# "):
            document.add_paragraph(line[2:], "Title" if first else "Heading 1")
            first = False
        elif line.startswith("## "):
            document.add_paragraph(line[3:], "Heading 2")
        elif line.startswith("- "):
            document.add_paragraph(line[2:], "List Bullet")
        elif re.match(r"\d+\. ", line):
            paragraph = document.add_paragraph(line)
            paragraph.paragraph_format.left_indent = Cm(0.45)
            paragraph.paragraph_format.first_line_indent = Cm(-0.45)
        elif line.startswith("`"):
            paragraph = document.add_paragraph()
            run = paragraph.add_run(line.strip("`"))
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            document.add_paragraph(line)
    output = source.with_suffix(".docx")
    document.save(output)
    print(output)


if __name__ == "__main__":
    for name in ("Manual_do_Jogador", "Guia_de_Respostas"):
        build(Path(__file__).parent / (name + ".md"))
